#!/usr/bin/env python3
"""Build Excel audit: DDS profile gaps for departments with source upload packs."""

from __future__ import annotations

import json
import subprocess
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "xlsx"
PHP = ROOT / "tools" / "dds_source_profile_audit.php"

STATUS_FILL = {
    "Completed": PatternFill("solid", fgColor="C6EFCE"),
    "Partial": PatternFill("solid", fgColor="FFEB9C"),
    "Pending": PatternFill("solid", fgColor="FFC7CE"),
}
STATUS_FONT = {
    "Completed": Font(color="006100"),
    "Partial": Font(color="9C5700"),
    "Pending": Font(color="9C0006"),
}


def style_header(cell, color="1F4E79"):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def apply_status(cell, status: str):
    cell.value = status
    cell.fill = STATUS_FILL.get(status, PatternFill())
    cell.font = STATUS_FONT.get(status, Font())
    cell.alignment = Alignment(horizontal="center", vertical="top")


def autosize(ws, max_w=40):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 10
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            width = min(max(width, min(len(val) + 2, max_w)), max_w)
        ws.column_dimensions[letter].width = width


def member_comment(m: dict) -> str:
    missing = m.get("missing") or []
    # Drop optional photo into its own note style
    content = [x for x in missing if not x.startswith("photo")]
    photo = [x for x in missing if x.startswith("photo")]
    lines = []
    if m.get("status") == "Pending":
        lines.append("PENDING (not received)")
    elif m.get("status") == "Partial":
        lines.append("PARTIAL (received but incomplete)")
    if content:
        lines.append("Missing: " + ", ".join(content))
    if photo:
        lines.append("Missing photo (optional)")
    if not m.get("in_source_pack"):
        lines.append("Not named in source upload pack (may still be on HRMS)")
    if m.get("source") == "CV-only":
        lines.append("On website via CV pack only (not in HRMS filter)")
    return "\n".join(f"• {ln}" for ln in lines)


def dept_comment(d: dict) -> str:
    sections = []
    pending = [m for m in d["members"] if m["status"] == "Pending"]
    partial = [m for m in d["members"] if m["status"] == "Partial"]
    no_photo = [m for m in d["members"] if not m.get("photo")]
    if pending:
        lines = []
        for m in pending:
            miss = [x for x in (m.get("missing") or []) if not x.startswith("photo")]
            lines.append(f"  • {m['name']} — " + (", ".join(miss) if miss else "qualifications, experience, publications"))
        sections.append("PENDING (not received)\n" + "\n".join(lines))
    if partial:
        lines = []
        for m in partial:
            miss = [x for x in (m.get("missing") or []) if not x.startswith("photo")]
            lines.append(f"  • {m['name']} — " + (", ".join(miss) if miss else "incomplete"))
        sections.append("PARTIAL (received but incomplete)\n" + "\n".join(lines))
    if no_photo:
        sections.append("MISSING PHOTO (optional)\n" + "\n".join(f"  • {m['name']}" for m in no_photo))
    if d.get("source_names_not_on_site"):
        sections.append(
            "IN SOURCE PACK BUT NOT FOUND ON SITE\n"
            + "\n".join(f"  • {n}" for n in d["source_names_not_on_site"])
        )
    return "\n\n".join(sections)


def main():
    proc = subprocess.run(
        ["php", str(PHP)],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if proc.returncode != 0:
        raise SystemExit(proc.stderr or proc.stdout or "PHP audit failed")
    data = json.loads(proc.stdout)

    wb = Workbook()
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws = wb.active
    ws.title = "Department Summary"
    headers = [
        "Source folder",
        "Department",
        "Total faculty",
        "Completed",
        "Partial",
        "Pending",
        "% Completed",
        "Department status",
        "Comments",
    ]
    for i, h in enumerate(headers, 1):
        style_header(ws.cell(1, i, h))
    ws.freeze_panes = "A2"

    r = 2
    totals = Counter()
    for d in data["departments"]:
        totals["total"] += d["total"]
        totals["completed"] += d["completed"]
        totals["partial"] += d["partial"]
        totals["pending"] += d["pending"]
        pct = (d["completed"] / d["total"]) if d["total"] else 0
        vals = [
            d["folder"],
            d["name"],
            d["total"],
            d["completed"],
            d["partial"],
            d["pending"],
            pct,
            d["department_status"],
            dept_comment(d),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = thin
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 9))
        apply_status(ws.cell(r, 8), d["department_status"])
        ws.cell(r, 7).number_format = "0.0%"
        lines = (vals[8] or "").count("\n") + 1
        ws.row_dimensions[r].height = min(240, max(36, 12 * lines + 8))
        r += 1

    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    ws.cell(r, 2, "Selected departments")
    ws.cell(r, 3, totals["total"])
    ws.cell(r, 4, totals["completed"])
    ws.cell(r, 5, totals["partial"])
    ws.cell(r, 6, totals["pending"])
    t = totals["total"] or 1
    ws.cell(r, 7, totals["completed"] / t).number_format = "0.0%"
    st = (
        "Completed"
        if totals["completed"] == totals["total"] and totals["total"]
        else ("Pending" if totals["completed"] == 0 else "Partial")
    )
    apply_status(ws.cell(r, 8), st)
    ws.cell(
        r,
        9,
        f"PENDING (not received): {totals['pending']}\n"
        f"PARTIAL (received but incomplete): {totals['partial']}\n"
        f"See Faculty Detail for per-person gaps",
    )
    for c in range(1, 10):
        ws.cell(r, c).fill = PatternFill("solid", fgColor="D6DCE4")
        ws.cell(r, c).border = thin
        ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=(c == 9))
    autosize(ws)
    ws.column_dimensions["I"].width = 72

    wd = wb.create_sheet("Faculty Detail")
    dh = [
        "Source folder",
        "Department",
        "Name",
        "Designation",
        "Listed in source pack",
        "Data source",
        "Status",
        "Photo",
        "Qualifications",
        "Experience",
        "Publications",
        "Profile slug",
        "Profile URL",
        "Comments",
    ]
    for i, h in enumerate(dh, 1):
        style_header(wd.cell(1, i, h), "833C0C")
    wd.freeze_panes = "A2"
    rr = 2
    for d in data["departments"]:
        for m in sorted(d["members"], key=lambda x: (x["status"] != "Pending", x["status"] != "Partial", x["name"])):
            url = f"https://dds.riphahpsh.edu.pk/faculty-profile?n={m['profile_slug']}"
            vals = [
                d["folder"],
                d["name"],
                m["name"],
                m["designation"],
                "Yes" if m.get("in_source_pack") else "No",
                m.get("source"),
                m["status"],
                "Yes" if m.get("photo") else "No",
                m.get("qualifications", 0),
                m.get("experience", 0),
                m.get("publications", 0),
                m.get("profile_slug"),
                url,
                member_comment(m),
            ]
            for c, v in enumerate(vals, 1):
                cell = wd.cell(rr, c, v)
                cell.border = thin
                cell.alignment = Alignment(vertical="top", wrap_text=(c == 14))
            apply_status(wd.cell(rr, 7), m["status"])
            rr += 1
    autosize(wd, 42)
    wd.column_dimensions["N"].width = 48

    wl = wb.create_sheet("Criteria")
    wl["A1"] = "DDS faculty profile gap audit (source upload folders)"
    wl["A1"].font = Font(bold=True, size=14, color="1F4E79")
    wl["A3"] = "Generated"
    wl["B3"] = data.get("generated_at")
    wl["A4"] = "Branch / commit"
    wl["B4"] = f"{data.get('branch')} @ {data.get('commit')}"
    wl["A5"] = "HRMS"
    wl["B5"] = "OK" if data.get("hrms_ok") else "FAILED"
    wl["A7"] = "Completed"
    wl["B7"] = data["criteria"]["Completed"]
    apply_status(wl["A7"], "Completed")
    wl["A8"] = "Partial"
    wl["B8"] = data["criteria"]["Partial"] + " — e.g. photo/some fields received, but still missing qualifications and/or experience/publications"
    apply_status(wl["A8"], "Partial")
    wl["A9"] = "Pending"
    wl["B9"] = data["criteria"]["Pending"]
    apply_status(wl["A9"], "Pending")
    wl["A11"] = "Department Completed / Partial / Pending"
    wl["B11"] = "All complete / some complete / none complete"
    wl["A13"] = "Source packs checked"
    wl["B13"] = (
        r"d:\New folder\pdc\Facilities required to be uploaded on website "
        r"+ Oral Pathology master/annexure docs "
        r"(Prosthodontics, Science of Dental Materials, Community & Preventive Dentistry, DDE, "
        r"Oral & Maxillofacial Surgery, Oral Biology, Oral Medicine, Paediatric Dentistry, "
        r"Periodontology, Oral Pathology)"
    )
    wl["A13"].alignment = Alignment(wrap_text=True)
    wl.merge_cells("B13:F15")
    autosize(wl, 80)

    OUT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    path = OUT / f"DDS_Faculty_Profile_Gaps_Source_Packs_{stamp}.xlsx"
    latest = OUT / "DDS_Faculty_Profile_Gaps_Source_Packs_Latest.xlsx"
    wb.save(path)
    try:
        wb.save(latest)
    except PermissionError:
        print("Latest locked; wrote timestamped only")
    print("Wrote", path)
    print("Wrote", latest)
    print(
        f"Faculty {totals['total']} | Completed={totals['completed']} Partial={totals['partial']} Pending={totals['pending']}"
    )
    for d in data["departments"]:
        print(
            f"{d['folder']}: {d['department_status']} "
            f"(done {d['completed']}/{d['total']}; pending {d['pending']}; partial {d['partial']})"
        )
        if d.get("source_names_not_on_site"):
            print("  NOT ON SITE:", ", ".join(d["source_names_not_on_site"]))


if __name__ == "__main__":
    main()
